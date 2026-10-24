#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Export the handbook to a printable Word document (概念手册.docx).

Usage:
    python scripts/export_handbook_docx.py [--out 概念手册.docx]
"""

import argparse
import json
import re
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SUBJECTS = [
    ("药理学", ROOT / "data/handbook-pharm.js"),
    ("有机化学", ROOT / "data/handbook-organic.js"),
    ("分析化学", ROOT / "data/handbook-analytical.js"),
]
WINDOW_RE = re.compile(r"window\.(\w+)\s*=\s*")
BODY_FONT = "微软雅黑"


def load(path):
    text = path.read_text(encoding="utf-8-sig")
    m = WINDOW_RE.search(text)
    if not m:
        raise SystemExit("[%s] missing window.XXX assignment" % path)
    body = text[m.end():].strip()
    if body.endswith(";"):
        body = body[:-1].strip()
    return json.loads(body)


def set_run_font(run, size=10.5, bold=False, color=None):
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        rfonts.set(qn("w:eastAsia"), BODY_FONT)


def style_heading(document, name, size, color=None, bold=True):
    st = document.styles[name]
    st.font.name = BODY_FONT
    st.font.size = Pt(size)
    st.font.bold = bold
    if color:
        st.font.color.rgb = RGBColor(*color)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)


def add_cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "概念手册.docx"))
    args = ap.parse_args()

    data = []
    for name, path in SUBJECTS:
        if not path.exists():
            raise SystemExit("missing %s, run build_handbook.py first" % path)
        data.append((name, load(path)))

    total_units = sum(len(d["chapters"]) for _, d in data)
    total_defs = sum(len(b.get("items") or []) for _, d in data for ch in d["chapters"] for b in ch["blocks"] if b.get("t") == "defs")
    total_points = sum(len(b.get("items") or []) for _, d in data for ch in d["chapters"] for b in ch["blocks"] if b.get("t") == "points")
    total_notes = sum(len(b.get("items") or []) for _, d in data for ch in d["chapters"] for b in ch["blocks"] if b.get("t") == "notes")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.25
    style_heading(doc, "Title", 24, (31, 41, 55))
    style_heading(doc, "Heading 1", 18, (29, 78, 216))
    style_heading(doc, "Heading 2", 14, (17, 24, 39))
    style_heading(doc, "Heading 3", 11.5, (37, 99, 235))

    title = doc.add_heading("考研药学综合 · 概念手册", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("三科 %d 单元 · 核心概念 %d 条 · 要点 %d 条 · 易混提示 %d 条" % (total_units, total_defs, total_points, total_notes))
    set_run_font(r, size=10, color=(107, 114, 128))
    tip = doc.add_paragraph()
    tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tip.add_run("按单元通读：先看核心概念，再过要点，最后看对比表与易混提示。打印版建议双面 A4。")
    set_run_font(r, size=9, color=(107, 114, 128))

    for si, (name, d) in enumerate(data):
        if si > 0:
            doc.add_page_break()
        doc.add_heading(name, level=1)
        for ch in d["chapters"]:
            doc.add_heading(ch["title"], level=2)
            if ch.get("summary"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run(ch["summary"])
                set_run_font(r, size=9.5, color=(107, 114, 128))
                r.italic = True
            for b in ch["blocks"]:
                t = b.get("t")
                doc.add_heading(b.get("title") or {"defs": "核心概念", "points": "要点梳理", "table": "对比表", "notes": "易混与记忆"}.get(t, ""), level=3)
                if t == "defs":
                    for it in b.get("items") or []:
                        p = doc.add_paragraph()
                        p.paragraph_format.space_after = Pt(2)
                        r = p.add_run(it.get("k") + "：")
                        set_run_font(r, bold=True)
                        r2 = p.add_run(it.get("v"))
                        set_run_font(r2)
                elif t == "points":
                    for i, it in enumerate(b.get("items") or [], 1):
                        p = doc.add_paragraph()
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.left_indent = Cm(0.6)
                        p.paragraph_format.first_line_indent = Cm(-0.6)
                        r = p.add_run("%d. %s" % (i, it))
                        set_run_font(r)
                elif t == "table":
                    head = b.get("head") or []
                    rows = b.get("rows") or []
                    table = doc.add_table(rows=len(rows) + 1, cols=len(head))
                    table.style = "Table Grid"
                    table.autofit = True
                    for ci, h in enumerate(head):
                        add_cell_text(table.rows[0].cells[ci], h, bold=True, size=9)
                    for ri, row in enumerate(rows, 1):
                        for ci, cell in enumerate(row):
                            add_cell_text(table.rows[ri].cells[ci], cell, size=9)
                    doc.add_paragraph().paragraph_format.space_after = Pt(0)
                elif t == "notes":
                    for it in b.get("items") or []:
                        p = doc.add_paragraph()
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.left_indent = Cm(0.35)
                        r = p.add_run("提示：")
                        set_run_font(r, bold=True, color=(180, 83, 9))
                        r2 = p.add_run(it)
                        set_run_font(r2)

    doc.save(args.out)
    print("wrote %s: %d units, %d defs, %d points" % (args.out, total_units, total_defs, total_points))


if __name__ == "__main__":
    main()
